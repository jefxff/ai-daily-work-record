"""把一条回答的引用导成 Excel。

为什么不复用 ``records/crud.make_export_view``：那个是"整表导出"（按模块
全量取数 + 行数上限 + 并发闸门），而这里的语义是"把刚才这次问答看到的几行
带走"——行数少、跨多个模块、还要带来源列。硬套会让两边的参数都变复杂。

安全两点沿用既有做法：
- 公式注入：``sanitize_xlsx_formulas`` 把 ``=`` 开头的单元格改成文本；
- 权限：只导出当前**仍有**读权限的模块，权限被收回的引用不出现。
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from work_record.export_utils import sanitize_xlsx_formulas

from . import catalog, registry, security

MAX_ROWS_PER_SHEET = 500
HEADER_FILL = PatternFill('solid', fgColor='F2F4F7')


def _cell_value(obj, field):
    display = getattr(obj, f'get_{field}_display', None)
    if callable(display):
        try:
            return str(display())
        except Exception:  # noqa: BLE001 - 选项字段取不到就退回原值
            pass
    value = getattr(obj, field, '')
    if value is None:
        return ''
    if isinstance(value, (list, tuple)):
        return '、'.join(str(v) for v in value)
    return value if isinstance(value, (int, float)) else str(value)


def _sheet_title(label, used):
    title = (label or '记录')[:28]
    candidate, index = title, 2
    while candidate in used:
        candidate = f'{title[:26]}-{index}'
        index += 1
    used.add(candidate)
    return candidate


def build_workbook(message, user, policy=None):
    """返回 ``Workbook``；没有任何可导出引用时返回 ``None``。"""
    policy = policy or security.policy()
    visible = {m.key for m in registry.available_modules(user, True, policy)}
    groups = {}
    for ref in message.refs or []:
        module = ref.get('module')
        if not module or module not in visible:
            continue          # 无权模块不导出（与历史回答复核同一口径）
        groups.setdefault(module, []).append(ref)
    if not groups:
        return None

    workbook = Workbook()
    workbook.remove(workbook.active)
    used_titles = set()

    for module_key, refs in groups.items():
        spec = catalog.get_module(module_key)
        if spec is None:
            continue
        model = catalog.model_of(spec)
        fields = [f for f in (spec.list_fields or (spec.title_field,)) if f]
        ids = [r.get('id') for r in refs if r.get('id') is not None]
        rows = []
        if ids:
            rows = list(model.objects.filter(pk__in=ids)[:MAX_ROWS_PER_SHEET])

        sheet = workbook.create_sheet(_sheet_title(spec.label, used_titles))
        headers = ['序号'] + [catalog.field_label(spec, f) for f in fields] + ['来源']
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(vertical='center')

        for index, obj in enumerate(rows, 1):
            sheet.append([index] + [_cell_value(obj, f) for f in fields] + [spec.label])

        if not rows:
            # 引用里的记录可能已被删除：给一行说明，别产出空 sheet 让人以为文件坏了
            sheet.append(['-'] + ['（该记录已删除或无权查看）'] + [''] * len(fields))

        widths = [6] + [max(14, min(36, len(str(catalog.field_label(spec, f))) * 3 + 8))
                        for f in fields] + [12]
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.freeze_panes = 'A2'

    if not workbook.sheetnames:
        return None
    return sanitize_xlsx_formulas(workbook)


def to_bytes(workbook):
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
