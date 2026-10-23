"""档案模块表格导入 / 导出业务引擎。

从 views.py 拆出的纯业务逻辑层：导入文件解析、列映射、字段写入 setter、
表级 upsert、Excel 导出响应构造等。视图层（views.py 及各域子模块）仅保留
请求/响应处理薄壳。函数实现与拆分前逐字节一致，仅模块位置变化。
"""

import csv
import io

# .xls 老格式解析依赖（openpyxl 只支持 .xlsx）。
# 采用"可选导入 + 明确报错"而不是模块级硬导入：缺依赖时应用仍能启动，
# 只是上传 .xls 时给出可读提示，而不是一个 ImportError 500。
# requirements.txt 已声明 xlrd==2.0.1（xlrd 2.x 只支持 .xls）。
try:
    import xlrd
except ImportError:  # pragma: no cover - 仅在依赖缺失的环境触发
    xlrd = None

from django.db import IntegrityError, transaction
from django.http import Http404, HttpResponse
from django.utils import timezone
from urllib.parse import quote

from ..models import (
    AnnualReview, Archive, BorrowRecord, Cabinet, Category, Ledger,
    License, LicenseBorrowRecord,
)
from ..validators import validate_import_file


def _read_import_rows(f):
    """解析上传的导入文件，返回字典行列表。支持 .xlsx / .xls / .csv。"""
    validate_import_file(f)
    name = (getattr(f, "name", "") or "").lower()
    if name.endswith(".csv"):
        raw = f.read()
        content = None
        for enc in ("utf-8-sig", "gbk"):
            try:
                content = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if content is None:
            raise ValueError("文件编码无法识别，请使用 UTF-8 编码 CSV")
        return list(csv.DictReader(io.StringIO(content)))
    if name.endswith(".xls"):
        if xlrd is None:
            raise ValueError(
                "服务器缺少 xlrd 依赖，暂时无法解析 .xls 文件；"
                "请让管理员执行 pip install -r requirements.txt，"
                "或把文件另存为 .xlsx 后重新上传。"
            )
        wb = xlrd.open_workbook(file_contents=f.read())
        ws = wb.sheet_by_index(0)
        if ws.nrows < 2:
            return []
        headers = [str(ws.cell_value(0, c) or "").strip() for c in range(ws.ncols)]
        out = []
        for r in range(1, ws.nrows):
            row = {}
            for c in range(ws.ncols):
                if c < len(headers) and headers[c]:
                    row[headers[c]] = str(ws.cell_value(r, c) or "").strip()
            if any(v for v in row.values()):
                out.append(row)
        return out
    # .xlsx
    from openpyxl import load_workbook
    wb = load_workbook(f, data_only=True)
    ws = wb["档案数据"] if "档案数据" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if all(v is None or str(v).strip() == "" for v in r):
            continue
        out.append({h: ("" if v is None else str(v).strip()) for h, v in zip(headers, r)})
    return out


# 智能导入引擎：字段映射表
FIELD_CHOICES = [
    ("code", "档案编号"),
    ("name", "档案名称/标题"),
    ("counterpart", "公司名称/相对方"),
    ("contract_no", "合同编号"),
    ("amount", "费用/金额"),
    ("sign_date", "签订日期"),
    ("expire_date", "截止/到期日期"),
    ("department", "部门"),
    ("security_level", "保密等级"),
    ("archive_label", "档案标签"),
    ("location_code", "位置编码"),
    ("handler", "移交人/经办人"),
    ("copies", "份数"),
    ("attachment_desc", "附件"),
    ("description", "备注"),
    ("tags", "标签"),
]


def _parse_excel_sheets(f):
    """解析 Excel 文件所有 sheet，返回 {sheet_name: {headers: [], rows: [{col: val}]}}。"""
    validate_import_file(f)
    name = (getattr(f, "name", "") or "").lower()
    result = {}
    if name.endswith(".xls"):
        if xlrd is None:
            raise ValueError(
                "服务器缺少 xlrd 依赖，暂时无法解析 .xls 文件；"
                "请让管理员执行 pip install -r requirements.txt，"
                "或把文件另存为 .xlsx 后重新上传。"
            )
        wb = xlrd.open_workbook(file_contents=f.read())
        for idx in range(wb.nsheets):
            ws = wb.sheet_by_index(idx)
            if ws.nrows < 2:
                continue
            headers = [str(ws.cell_value(0, c) or "").strip() for c in range(ws.ncols)]
            rows = []
            for r in range(1, ws.nrows):
                row = {}
                for c in range(ws.ncols):
                    if c < len(headers) and headers[c]:
                        row[headers[c]] = str(ws.cell_value(r, c) or "").strip()
                if any(row.values()):
                    rows.append(row)
            if rows:
                result[ws.name] = {"headers": [h for h in headers if h], "rows": rows}
        return result
    from openpyxl import load_workbook
    wb = load_workbook(f, data_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            continue
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        data_rows = []
        for r in rows[1:]:
            row = {}
            for h, v in zip(headers, r):
                if h:
                    row[h] = ("" if v is None else str(v).strip())
            if any(row.values()):
                data_rows.append(row)
        if data_rows:
            result[ws.title] = {"headers": [h for h in headers if h], "rows": data_rows}
    return result


def _auto_map_columns(headers):
    """根据表头文本自动推断系统字段映射。返回 {col_name: field_name}。"""
    mapping = {}
    keyword_map = {
        "code": ["档案编号", "编号", "序列", "序号", "序 号"],
        "name": ["标题", "名称", "档案名称", "文件名称", "合同名称"],
        "counterpart": ["公司名称", "相对方", "相对人", "供应商", "合作方", "发文单位", "签订单位", "名称/姓名", "单位", "公司名称/相对人"],
        "contract_no": ["合同编号", "单号", "编号", "文书编号", "文件编号", "报告编号"],
        "amount": ["费用", "金额", "额度", "保费"],
        "sign_date": ["签订日期", "签发日期", "公证日期", "申请日期", "日期", "签订时间"],
        "expire_date": ["到期日", "截止日期", "截止日期", "到期时间", "截止时间"],
        "department": ["部门"],
        "security_level": ["保密等级"],
        "archive_label": ["档案名称", "档案标签"],
        "location_code": ["位置", "归位处"],
        "handler": ["移交人", "经办人", "经手人", "移送人"],
        "copies": ["份数"],
        "attachment_desc": ["附件"],
        "description": ["备注"],
        "tags": ["标签"],
    }
    for col in headers:
        for field, keywords in keyword_map.items():
            if col.strip() in keywords and field not in mapping.values():
                mapping[col] = field
                break
    return mapping


def _csv_safe(v):
    """CSV 导出防公式注入：'='/'+'/'-'/'@'/Tab 开头的文本会被 Excel 当公式解析。"""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


# ---------- 表级通用导入 / 导出（导出 → Excel 修改 → 导入更新 闭环） ----------

def _parse_date_val(raw):
    """多格式日期解析，失败返回 None。"""
    from datetime import datetime as _dt
    raw = raw.split(" ")[0] if " " in raw else raw
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%m/%d/%Y"):
        try:
            return _dt.strptime(raw, fmt).date()
        except ValueError:
            continue
    return None


def _parse_datetime_val(raw):
    """多格式日期时间解析，失败返回 None（naive 值转为当前时区 aware）。"""
    from datetime import datetime as _dt
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%Y/%m/%d %H:%M", "%Y/%m/%d"):
        try:
            v = _dt.strptime(raw, fmt)
            if timezone.is_naive(v):
                v = timezone.make_aware(v)
            return v
        except ValueError:
            continue
    return None


def _fmt_date(v):
    return v.strftime("%Y-%m-%d") if v else ""


def _fmt_datetime(v):
    if not v:
        return ""
    if timezone.is_aware(v):
        v = timezone.localtime(v)
    return v.strftime("%Y-%m-%d %H:%M")


def _category_from_path(path):
    """按「A/B/C」全路径逐级 get_or_create 类别，返回末级节点；空路径返回 None。"""
    node = None
    for part in [p.strip() for p in path.split("/") if p.strip()]:
        node, _ = Category.objects.get_or_create(name=part[:100], parent=node)
    return node


def _col(header, getter, setter=None):
    """列定义：(表头, 取值函数, 写值函数)。setter 为 None 表示该列仅导出展示。"""
    return (header, getter, setter)


def _setter_int(field):
    def _s(obj, raw):
        try:
            setattr(obj, field, int(float(raw)))
        except (TypeError, ValueError):
            raise ValueError(f"{field} 需为数字，收到「{raw}」")
    return _s


def _setter_date(field):
    def _s(obj, raw):
        v = _parse_date_val(raw)
        if v is None:
            raise ValueError(f"日期格式无法识别：{raw}（支持 2026-01-31 / 2026/1/31）")
        setattr(obj, field, v)
    return _s


def _setter_datetime(field):
    def _s(obj, raw):
        v = _parse_datetime_val(raw)
        if v is None:
            raise ValueError(f"时间格式无法识别：{raw}")
        setattr(obj, field, v)
    return _s


def _setter_text(field, max_len=200):
    def _s(obj, raw):
        setattr(obj, field, raw[:max_len] if field != "description" and field != "remark" else raw)
    return _s


def _setter_fk_ledger(obj, raw):
    obj.ledger, _ = Ledger.objects.get_or_create(name=raw[:200])


def _setter_fk_category(obj, raw):
    obj.category = _category_from_path(raw)


def _setter_fk_cabinet(obj, raw):
    cab = Cabinet.objects.filter(code=raw).first()
    if not cab:
        raise ValueError(f"档案柜「{raw}」不存在，请先创建")
    obj.cabinet = cab


def _setter_fk_archive(obj, raw):
    arc = Archive.objects.filter(code=raw).first()
    if not arc:
        raise ValueError(f"档案编号「{raw}」不存在")
    obj.archive = arc


def _setter_fk_license(obj, raw):
    lic = License.objects.filter(code=raw).first()
    if not lic:
        raise ValueError(f"证照编号「{raw}」不存在")
    obj.license = lic


def _setter_cabinet_layer_cells(obj, raw):
    """「各层格数」列 setter：逗号分隔数字（如 10,10,8,8），暂存到实例待保存后同步。"""
    parts = [p.strip() for p in str(raw).replace("，", ",").split(",") if p.strip()]
    if not parts:
        raise ValueError("各层格数不能为空，格式如 10,10,8,8")
    try:
        obj._pending_layer_cells = [int(float(p)) for p in parts]
    except (TypeError, ValueError):
        raise ValueError(f"各层格数格式无法识别：{raw}（应为逗号分隔数字，如 10,10,8,8）")


def _finalize_cabinet(obj):
    """Cabinet 保存后同步层记录（导入引擎 / 表单路径共用）。"""
    if isinstance(obj, Cabinet):
        obj.sync_layers(getattr(obj, "_pending_layer_cells", None))
        if hasattr(obj, "_pending_layer_cells"):
            del obj._pending_layer_cells


TABLE_SPECS = {
    "archives": {
        "label": "档案",
        "model": Archive,
        "match_field": "code",
        "match_header": "编号",
        "redirect": "archive:overview",
        "queryset": lambda: Archive.objects.select_related("cabinet", "category", "ledger").order_by("id"),
        "columns": [
            _col("编号", lambda a: a.code),
            _col("名称", lambda a: a.name, _setter_text("name", 200)),
            _col("所属台账", lambda a: a.ledger.name if a.ledger_id else "", _setter_fk_ledger),
            _col("类别", lambda a: a.category.full_path if a.category_id else "", _setter_fk_category),
            _col("档案柜", lambda a: a.cabinet.code if a.cabinet_id else "", _setter_fk_cabinet),
            _col("层", lambda a: a.layer if a.layer is not None else "", _setter_int("layer")),
            _col("格", lambda a: a.cell if a.cell is not None else "", _setter_int("cell")),
            _col("位置编码", lambda a: a.location_code, _setter_text("location_code", 50)),
            _col("来源Sheet", lambda a: a.source_sheet, _setter_text("source_sheet", 100)),
            _col("公司名称/相对方", lambda a: a.counterpart, _setter_text("counterpart")),
            _col("合同编号", lambda a: a.contract_no, _setter_text("contract_no", 100)),
            _col("金额", lambda a: a.amount, _setter_text("amount", 50)),
            _col("签订日期", lambda a: _fmt_date(a.sign_date), _setter_date("sign_date")),
            _col("到期日期", lambda a: _fmt_date(a.expire_date), _setter_date("expire_date")),
            _col("部门", lambda a: a.department, _setter_text("department", 50)),
            _col("保密等级", lambda a: a.security_level, _setter_text("security_level", 20)),
            _col("档案标签", lambda a: a.archive_label, _setter_text("archive_label", 100)),
            _col("经办人", lambda a: a.handler, _setter_text("handler", 50)),
            _col("份数", lambda a: a.copies if a.copies is not None else "", _setter_int("copies")),
            _col("附件描述", lambda a: a.attachment_desc, _setter_text("attachment_desc")),
            _col("标签", lambda a: a.tags, _setter_text("tags")),
            _col("档案描述", lambda a: a.description, _setter_text("description")),
        ],
    },
    "ledgers": {
        "label": "台账",
        "model": Ledger,
        "match_field": "name",
        "match_header": "台账名称",
        "redirect": "archive:overview",
        "queryset": lambda: Ledger.objects.order_by("id"),
        "columns": [
            _col("台账名称", lambda l: l.name),
            _col("说明", lambda l: l.description, _setter_text("description", 300)),
            _col("排序", lambda l: l.sort_order, _setter_int("sort_order")),
        ],
    },
    "cabinets": {
        "label": "档案柜",
        "model": Cabinet,
        "match_field": "code",
        "match_header": "档案柜编号",
        "redirect": "archive:overview",
        "queryset": lambda: Cabinet.objects.prefetch_related("layer_configs").order_by("id"),
        "columns": [
            _col("档案柜编号", lambda c: c.code),
            _col("所在位置", lambda c: c.location, _setter_text("location")),
            _col("层数", lambda c: c.layers, _setter_int("layers")),
            _col("各层格数", lambda c: c.layer_cells_summary(), _setter_cabinet_layer_cells),
        ],
    },
    "categories": {
        "label": "档案类别",
        "model": Category,
        "match_field": "path",
        "match_header": "类别全路径",
        "redirect": "archive:category_list",
        "queryset": lambda: Category.objects.order_by("id"),
        "columns": [
            _col("类别全路径", lambda c: c.full_path),
            _col("排序", lambda c: c.sort_order, _setter_int("sort_order")),
        ],
    },
    "borrows": {
        "label": "借阅记录",
        "model": BorrowRecord,
        "match_field": "id",
        "match_header": "记录ID",
        "redirect": "archive:borrow_list",
        "queryset": lambda: BorrowRecord.objects.select_related("archive").order_by("id"),
        "columns": [
            _col("记录ID", lambda r: r.pk),
            _col("档案编号", lambda r: r.archive.code, _setter_fk_archive),
            _col("借阅人", lambda r: r.borrower, _setter_text("borrower", 50)),
            _col("借阅时间", lambda r: _fmt_datetime(r.borrow_time), _setter_datetime("borrow_time")),
            _col("应归还日期", lambda r: _fmt_date(r.due_time), _setter_date("due_time")),
            _col("实际归还时间", lambda r: _fmt_datetime(r.return_time), _setter_datetime("return_time")),
            _col("状态", lambda r: r.get_status_display()),
            _col("备注", lambda r: r.remark, _setter_text("remark")),
        ],
    },
    "licenses": {
        "label": "证照",
        "model": License,
        "match_field": "code",
        "match_header": "证照编号",
        "redirect": "archive:license_list",
        "queryset": lambda: License.objects.order_by("id"),
        "columns": [
            _col("证照编号", lambda l: l.code),
            _col("证照名称", lambda l: l.name, _setter_text("name")),
            _col("持证单位", lambda l: l.holder, _setter_text("holder")),
            _col("发证机关", lambda l: l.issuer, _setter_text("issuer")),
            _col("发证日期", lambda l: _fmt_date(l.issue_date), _setter_date("issue_date")),
            _col("有效期至", lambda l: _fmt_date(l.expire_date), _setter_date("expire_date")),
            _col("年审周期(月)", lambda l: l.review_cycle, _setter_int("review_cycle")),
            _col("下次年审日", lambda l: _fmt_date(l.next_review_date), _setter_date("next_review_date")),
            _col("状态", lambda l: l.status_display),
            _col("备注", lambda l: l.remark, _setter_text("remark")),
        ],
    },
    "reviews": {
        "label": "年审记录",
        "model": AnnualReview,
        "match_field": "id",
        "match_header": "记录ID",
        "redirect": "archive:license_list",
        "queryset": lambda: AnnualReview.objects.select_related("license").order_by("id"),
        "columns": [
            _col("记录ID", lambda r: r.pk),
            _col("证照编号", lambda r: r.license.code, _setter_fk_license),
            _col("年审日期", lambda r: _fmt_date(r.review_date), _setter_date("review_date")),
            _col("年审结果", lambda r: r.get_result_display(), None),
            _col("经办人", lambda r: r.reviewer, _setter_text("reviewer", 50)),
            _col("新有效期至", lambda r: _fmt_date(r.new_expire_date), _setter_date("new_expire_date")),
            _col("备注", lambda r: r.remark, _setter_text("remark", 500)),
        ],
    },
    "license_borrows": {
        "label": "证照借阅记录",
        "model": LicenseBorrowRecord,
        "match_field": "id",
        "match_header": "记录ID",
        "redirect": "archive:license_borrow_list",
        "queryset": lambda: LicenseBorrowRecord.objects.select_related("license").order_by("id"),
        "columns": [
            _col("记录ID", lambda r: r.pk),
            _col("证照编号", lambda r: r.license.code, _setter_fk_license),
            _col("借阅人", lambda r: r.borrower, _setter_text("borrower", 50)),
            _col("借阅时间", lambda r: _fmt_datetime(r.borrow_time), _setter_datetime("borrow_time")),
            _col("应归还日期", lambda r: _fmt_date(r.due_time), _setter_date("due_time")),
            _col("实际归还时间", lambda r: _fmt_datetime(r.return_time), _setter_datetime("return_time")),
            _col("状态", lambda r: r.get_status_display()),
            _col("备注", lambda r: r.remark, _setter_text("remark")),
        ],
    },
}

# 年审结果中文 ↔ 代码映射（导入解析用）
_REVIEW_RESULT_MAP = {v: k for k, v in dict(AnnualReview.RESULT_CHOICES).items()}

# 档案列定义快捷引用（台账维度导入 / 导出视图共用）
_ARCHIVE_COLS = TABLE_SPECS["archives"]["columns"]
_ARCHIVE_MATCH_HEADER = TABLE_SPECS["archives"]["match_header"]  # 「编号」


def _get_spec(table):
    spec = TABLE_SPECS.get(table)
    if not spec:
        raise Http404("未知的导入导出对象")
    return spec


def _match_obj(spec, key_raw):
    """按匹配键定位现有记录；未命中返回 None。"""
    if not key_raw:
        return None
    model = spec["model"]
    field = spec["match_field"]
    if field == "id":
        return model.objects.filter(pk=key_raw).first() if str(key_raw).isdigit() else None
    if field == "path":
        return None  # 类别按路径新建/定位，不走精确匹配
    return model.objects.filter(**{field: key_raw}).first()


def _apply_columns(spec, obj, row, is_created, field_errors=None):
    """按列定义把行数据写入对象；空值跳过（保留原值）。返回是否写入了任何字段。

    field_errors 传列表时启用列级容错：某列写值失败（如档案柜编号不存在、日期格式错误）
    则跳过该列保留原值，错误记入 [(列名, 原因)]；不传时保持原语义（抛异常回滚整行）。
    """
    changed = False
    for header, _getter, setter in spec["columns"]:
        if setter is None or header == spec["match_header"]:
            continue
        raw = str(row.get(header, "") or "").strip()
        if not raw:
            continue
        if spec["label"] == "年审记录" and header == "年审结果":
            code = _REVIEW_RESULT_MAP.get(raw)
            if not code:
                raise ValueError(f"年审结果需为：通过/未通过/待处理，收到「{raw}」")
            obj.result = code
            changed = True
            continue
        try:
            setter(obj, raw)
            changed = True
        except Exception as e:
            if field_errors is None:
                raise
            field_errors.append((header, str(e)))
    # 记录表新建时给默认借阅时间/状态，避免必填约束直接崩
    if is_created and hasattr(obj, "borrow_time") and obj.borrow_time is None:
        obj.borrow_time = timezone.now()
    return changed


def _validate_archive_row(row):
    """档案行只读预检（不写库），返回将忽略的单元格 [(列名, 原因)]。"""
    errs = []

    def raw(h):
        return str(row.get(h, "") or "").strip()

    cab = raw("档案柜")
    if cab and not Cabinet.objects.filter(code=cab).exists():
        names = "、".join(Cabinet.objects.values_list("code", flat=True)[:8]) or "暂无"
        errs.append(("档案柜", f"编号「{cab}」不存在（系统现有：{names}），该列将保留原值"))
    for h in ("层", "格", "份数"):
        v = raw(h)
        if v:
            try:
                int(float(v))
            except (TypeError, ValueError):
                errs.append((h, f"「{v}」需为数字，该列将保留原值"))
    for h in ("签订日期", "到期日期"):
        v = raw(h)
        if v and _parse_date_val(v) is None:
            errs.append((h, f"「{v}」格式无法识别（支持 2026-01-31 / 2026/1/31），该列将保留原值"))
    return errs


def _export_response(spec, include_header_only=False):
    """生成 xlsx 导出响应（table_export / table_template 共用）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = spec["label"][:30]
    headers = [c[0] for c in spec["columns"]]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E6F7FF")
    if not include_header_only:
        for obj in spec["queryset"]():
            ws.append([c[1](obj) for c in spec["columns"]])
    from work_record.export_utils import sanitize_xlsx_formulas
    sanitize_xlsx_formulas(wb)   # 防公式注入（审计 #27）
    buf = io.BytesIO()
    wb.save(buf)
    suffix = "模板" if include_header_only else "导出"
    fname = f"{spec['label']}{suffix}_{timezone.localtime():%Y%m%d_%H%M}.xlsx"
    resp = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    from work_record.export_utils import content_disposition
    resp["Content-Disposition"] = content_disposition(fname)   # 中文名 + 头注入防护
    return resp


# ---------- 台账档案导入 / 导出（按台账维度）引擎 ----------

def _archive_book_bytes(sheets):
    """sheets: [(工作表名, 档案列表)] → xlsx 二进制。列结构与档案导出一致（一个台账一个工作表）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    wb.remove(wb.active)
    for name, archives in sheets:
        ws = wb.create_sheet(title=(name or "档案")[:31])
        ws.append([c[0] for c in _ARCHIVE_COLS])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E6F7FF")
        for a in archives:
            ws.append([c[1](a) for c in _ARCHIVE_COLS])
    from work_record.export_utils import sanitize_xlsx_formulas
    sanitize_xlsx_formulas(wb)   # 防公式注入（审计 #27）
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_response(data, fname):
    resp = HttpResponse(
        data,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    from work_record.export_utils import content_disposition
    resp["Content-Disposition"] = content_disposition(fname)
    return resp


def _classify_archive_rows(rows):
    """按「编号」判定行的新增/更新。返回 (keys, created_n, updated_n, existing_keys)。"""
    keys = [str(r.get(_ARCHIVE_MATCH_HEADER, "") or "").strip() for r in rows]
    existing = set()
    vals = [k for k in keys if k]
    for i in range(0, len(vals), 500):
        batch = vals[i:i + 500]
        existing.update(Archive.objects.filter(code__in=batch).values_list("code", flat=True))
    created_n = sum(1 for k in keys if k not in existing)
    return keys, created_n, len(keys) - created_n, existing


def _upsert_archive_row(row, ledger):
    """单行 upsert：按「编号」匹配更新（仅非空列，无效单元格保留原值），不存在则新建并归入指定台账。

    返回 (是否新建, 忽略的无效单元格 [(列名, 原因)])。
    """
    key_raw = str(row.get(_ARCHIVE_MATCH_HEADER, "") or "").strip()
    field_errors = []
    try:
        with transaction.atomic():
            obj = Archive.objects.filter(code=key_raw).first() if key_raw else None
            created = obj is None
            if created:
                if not key_raw:
                    raise ValueError("「编号」为空，无法定位或新建")
                obj = Archive(code=key_raw[:100], ledger=ledger)
            _apply_columns(TABLE_SPECS["archives"], obj, row, is_created=created, field_errors=field_errors)
            obj.ledger = ledger  # 归属以操作维度为准，不信任 Excel 中的台账列
            if hasattr(obj, "refresh_status"):
                obj.refresh_status()
            obj.save()
    except IntegrityError:
        # 同一柜/层/格重复写入：DB 唯一约束兜底（评审 R1），逐行回滚不影响其他行
        raise ValueError("该存放位置（柜/层/格）已被其他档案占用，已跳过")
    return created, field_errors
