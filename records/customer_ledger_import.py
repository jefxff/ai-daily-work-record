"""客户台账 Excel 导入 / 导出引擎。

首期导入只接受 Excel（.xlsx / .xls）；来源表 31 列与导出表头完全一致，
保证「导出 → 修改 → 导入」闭环。附件不参与 Excel 导入：营业执照、法人身份证、
基本合同等扫描件在记录详情/编辑页合并为一个 PDF 上传。

去重优先级：统一社会信用代码 → 档案编号 → 客户名称+法人姓名。
默认策略：仅补空（不覆盖已有值）；向导中可切换为覆盖。
事务语义：预览阶段发现任何阻断错误即整批不写入；确认阶段单事务全成全败。
"""

import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.http import HttpResponse
from django.utils import timezone

from .models import CustomerLedger

MAX_IMPORT_ROWS = 5000
IMPORT_MAX_SIZE = 10 * 1024 * 1024

# 来源表 31 列（顺序即导出顺序，也是模板表头）
CUSTOMER_LEDGER_HEADERS = [
    '序号', '责任业务', '客户名称', '客户分类', '法人姓名', '地址',
    '统一社会信用代码', '注册资金', '实缴资金', '公司现状', '营业执照', '账期',
    '额度', '变更时间', '基本合同', '基本合同签订日期', '基本合同存在的问题',
    '资产担保', '法代身份证', '担保人身份证', '资产担保人', '担保资产资产明细',
    '担保资产是否抵押', '资产担保存在的问题', '合同份数', '处理建议', '系统信息',
    '位置', '档案编号', '归档日期', '备注',
]

# 纯文本列：Excel 表头 -> 模型字段
TEXT_FIELDS = {
    '责任业务': 'responsible_salesperson',
    '客户名称': 'customer_name',
    '客户分类': 'customer_category',
    '法人姓名': 'legal_representative',
    '地址': 'address',
    '统一社会信用代码': 'credit_code',
    '公司现状': 'company_status',
    '营业执照': 'business_license',
    '账期': 'payment_terms',
    '基本合同': 'basic_contract',
    '基本合同存在的问题': 'contract_issues',
    '资产担保': 'asset_guarantee',
    '法代身份证': 'legal_id_card',
    '担保人身份证': 'guarantor_id_card',
    '资产担保人': 'asset_guarantor',
    '担保资产资产明细': 'guaranteed_asset_detail',
    '担保资产是否抵押': 'asset_mortgaged',
    '资产担保存在的问题': 'asset_guarantee_issues',
    '合同份数': 'contract_count',
    '处理建议': 'handling_suggestion',
    '系统信息': 'system_info',
    '位置': 'location',
    '档案编号': 'archive_no',
    '备注': 'remarks',
}

# 金额列：Excel 表头 -> (数值字段, 原文保留字段)
AMOUNT_FIELDS = {
    '注册资金': ('registered_capital', 'registered_capital_text'),
    '实缴资金': ('paid_in_capital', 'paid_in_capital_text'),
    '额度': ('credit_limit', 'credit_limit_text'),
}

# 日期列：Excel 表头 -> 模型字段
DATE_FIELDS = {
    '变更时间': 'change_date',
    '基本合同签订日期': 'contract_sign_date',
    '归档日期': 'archive_date',
}

# 模型字段 -> 中文标签（预览变更清单 / 导出用）
FIELD_LABELS = {}
for _header, _field in TEXT_FIELDS.items():
    FIELD_LABELS[_field] = _header
for _header, (_num, _text) in AMOUNT_FIELDS.items():
    FIELD_LABELS[_num] = f'{_header}(万元)'
    FIELD_LABELS[_text] = f'{_header}原文'
for _header, _field in DATE_FIELDS.items():
    FIELD_LABELS[_field] = _header
FIELD_LABELS['contract_count'] = '合同份数'
FIELD_LABELS['archive_no'] = '档案编号'

# 导入可写字段（不含附件、创建/更新时间、序号）
IMPORT_FIELDS = [
    'responsible_salesperson', 'customer_name', 'customer_category',
    'legal_representative', 'address', 'credit_code',
    'registered_capital', 'registered_capital_text',
    'paid_in_capital', 'paid_in_capital_text', 'company_status',
    'business_license', 'payment_terms',
    'credit_limit', 'credit_limit_text', 'change_date',
    'basic_contract', 'contract_sign_date', 'contract_issues',
    'asset_guarantee', 'legal_id_card', 'guarantor_id_card',
    'asset_guarantor', 'guaranteed_asset_detail', 'asset_mortgaged',
    'asset_guarantee_issues', 'contract_count',
    'handling_suggestion', 'system_info', 'location',
    'archive_no', 'archive_date', 'remarks',
]

_REQUIRED_HEADER = '客户名称'


# ---------------------------------------------------------------- 单元格解析

def _cell_text(value):
    """任意单元格值 -> 去空白字符串（None -> ''，日期按 YYYY-MM-DD）。"""
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, date):
        return value.strftime('%Y-%m-%d')
    return str(value).strip()


def _to_decimal(raw):
    """金额原文 -> 万元 Decimal；无法结构化时返回 None（原文仍会保留）。

    支持 '50万'、'300万'、'0'、'50.5'、'500000元'；'/'、'未披露'、'查不到'
    等非数值一律返回 None。
    """
    text = _cell_text(raw)
    if not text or text in {'/', '-', '—', '无'}:
        return None
    lowered = text.replace(' ', '')
    if any(word in lowered for word in ('未披露', '查不到', '不详', '未知')):
        return None
    multiplier = Decimal('1')
    if lowered.endswith('万元'):
        lowered = lowered[:-2]
    elif lowered.endswith('万'):
        lowered = lowered[:-1]
    elif lowered.endswith('元'):
        lowered = lowered[:-1]
        multiplier = Decimal('0.0001')
    try:
        return (Decimal(lowered) * multiplier).quantize(Decimal('0.01'))
    except (InvalidOperation, ValueError):
        return None


def _to_date(raw):
    """日期单元格 -> date；无法识别返回 None。

    openpyxl/xlrd 会把真正的日期单元格解析为 datetime/date；文本日期兼容
    YYYY-MM-DD / YYYY/M/D / YYYY.M.D / YYYYMMDD。
    """
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = _cell_text(raw)
    if not text:
        return None
    text = text.split(' ')[0]
    for fmt in ('%Y-%m-%d', '%Y/%m/%d', '%Y.%m.%d', '%Y%m%d'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _empty(value):
    return value is None or value == ''


# ---------------------------------------------------------------- 文件解析

def _validate_file(f):
    from work_record.file_validation import content_matches, normalize_ext, read_head

    ext = normalize_ext(getattr(f, 'name', '') or '')
    if ext not in {'XLSX', 'XLS'}:
        raise ValueError(
            f'不支持的文件类型「{ext or "无扩展名"}」；客户台账导入仅支持 .xlsx / .xls'
        )
    if getattr(f, 'size', 0) > IMPORT_MAX_SIZE:
        raise ValueError(
            f'文件大小 {f.size / 1024 / 1024:.1f}MB 超过上限 '
            f'{IMPORT_MAX_SIZE // 1024 // 1024}MB'
        )
    if not content_matches(f, ext, read_head(f)):
        raise ValueError('文件内容与 Excel 扩展名不符，已拒绝解析')
    return ext


def _read_xlsx(f):
    from openpyxl import load_workbook

    wb = load_workbook(f, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return [], []
    headers = [_cell_text(h) for h in rows[0]]
    data_rows = []
    for raw in rows[1:]:
        if all(v is None or _cell_text(v) == '' for v in raw):
            continue
        data_rows.append({
            header: _cell_text(value)
            for header, value in zip(headers, raw) if header
        })
    return headers, data_rows


def _read_xls(f):
    try:
        import xlrd
    except ImportError:
        raise ValueError('服务器缺少 xlrd 依赖，暂时无法解析 .xls；请另存为 .xlsx 后重试')

    wb = xlrd.open_workbook(file_contents=f.read())
    sheet = wb.sheet_by_index(0)
    if sheet.nrows < 2:
        return [], []
    headers = [_cell_text(sheet.cell_value(0, c)) for c in range(sheet.ncols)]
    data_rows = []
    for r in range(1, sheet.nrows):
        row = {}
        for c in range(sheet.ncols):
            header = headers[c]
            if not header:
                continue
            cell = sheet.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    value = xlrd.xldate_as_datetime(cell.value, wb.datemode)
                    row[header] = _cell_text(value)
                except (ValueError, OverflowError):
                    row[header] = _cell_text(cell.value)
            else:
                row[header] = _cell_text(cell.value)
        if any(v for v in row.values()):
            data_rows.append(row)
    return headers, data_rows


def parse_excel(f):
    """解析 Excel，返回 [{表头: 文本}]；表头缺失/行数超限时抛 ValueError。"""
    ext = _validate_file(f)
    try:
        f.seek(0)
    except (AttributeError, OSError):
        pass
    headers, rows = _read_xls(f) if ext == 'XLS' else _read_xlsx(f)
    headers = [h for h in headers if h]
    if _REQUIRED_HEADER not in headers:
        raise ValueError('Excel 缺少必填列「客户名称」，请使用系统提供的导入模板')
    if len(rows) > MAX_IMPORT_ROWS:
        raise ValueError(f'单次导入最多 {MAX_IMPORT_ROWS} 行，当前 {len(rows)} 行')
    return rows


# ---------------------------------------------------------------- 行 -> 字段

def row_to_values(row):
    """Excel 行 -> 模型字段字典（含金额原文与结构化值、日期）。"""
    values = {}
    for header, field in TEXT_FIELDS.items():
        values[field] = _cell_text(row.get(header))
    for header, (num_field, text_field) in AMOUNT_FIELDS.items():
        raw = _cell_text(row.get(header))
        values[text_field] = raw
        values[num_field] = _to_decimal(raw)
    for header, field in DATE_FIELDS.items():
        values[field] = _to_date(row.get(header))
    # 档案编号空串必须落 None，避免唯一约束冲突
    values['archive_no'] = values.get('archive_no') or None
    return values


def _match_existing(values):
    """按「信用代码 → 档案编号 → 客户名称+法人」顺序定位已有记录。"""
    credit = values.get('credit_code')
    if credit:
        obj = CustomerLedger.objects.filter(credit_code=credit).first()
        if obj:
            return obj, '统一社会信用代码'
    archive_no = values.get('archive_no')
    if archive_no:
        obj = CustomerLedger.objects.filter(archive_no=archive_no).first()
        if obj:
            return obj, '档案编号'
    name = values.get('customer_name')
    legal = values.get('legal_representative')
    if name:
        qs = CustomerLedger.objects.filter(customer_name=name)
        if legal:
            qs = qs.filter(legal_representative=legal)
        obj = qs.first()
        if obj:
            return obj, '客户名称+法人'
    return None, ''


def _display(value):
    if value is None:
        return ''
    if isinstance(value, Decimal):
        return format(value.normalize(), 'f')
    if isinstance(value, (date, datetime)):
        return value.strftime('%Y-%m-%d')
    return str(value)


def _changes(values, obj, overwrite):
    """计算本行将写入的字段 [(字段标签, 新值)]；obj 为 None 表示新建。"""
    changes = []
    for field in IMPORT_FIELDS:
        new = values.get(field)
        if _empty(new):
            continue
        if obj is not None:
            old = getattr(obj, field)
            if not overwrite and not _empty(old):
                continue
            if old == new:
                continue
        changes.append((FIELD_LABELS.get(field, field), _display(new)))
    return changes


# ---------------------------------------------------------------- 预览

def build_preview(rows, overwrite=False):
    """构建导入预览：items / errors / warnings / summary。

    errors 为阻断错误（确认阶段整批拒绝）；warnings 为不影响导入的提示。
    """
    items, errors, warnings = [], [], []
    seen_credit, seen_archive = {}, {}

    for index, row in enumerate(rows, start=2):
        values = row_to_values(row)
        name = values.get('customer_name') or ''
        if not name:
            errors.append(f'第 {index} 行：客户名称为空')
            continue

        credit = values.get('credit_code')
        if credit:
            if credit in seen_credit:
                errors.append(
                    f'第 {index} 行：统一社会信用代码「{credit}」与第 {seen_credit[credit]} 行重复'
                )
                continue
            seen_credit[credit] = index

        archive_no = values.get('archive_no')
        if archive_no:
            if archive_no in seen_archive:
                errors.append(
                    f'第 {index} 行：档案编号「{archive_no}」与第 {seen_archive[archive_no]} 行重复'
                )
                continue
            seen_archive[archive_no] = index

        obj, matched_by = _match_existing(values)

        if archive_no and CustomerLedger.objects.filter(archive_no=archive_no).exclude(
            pk=getattr(obj, 'pk', 0) or 0
        ).exists():
            errors.append(f'第 {index} 行：档案编号「{archive_no}」已被其他记录占用')
            continue

        for header, (num_field, text_field) in AMOUNT_FIELDS.items():
            raw = values.get(text_field) or ''
            if raw and values.get(num_field) is None:
                warnings.append(
                    f'第 {index} 行：{header}「{raw}」无法转为数值，已保留原文'
                )

        changes = _changes(values, obj, overwrite)
        if obj is None:
            action = 'create'
        elif changes:
            action = 'update'
        else:
            action = 'skip'
        items.append({
            'row': index,
            'action': action,
            'values': values,
            'customer_name': name,
            'matched_by': matched_by,
            'matched_id': getattr(obj, 'pk', None),
            'matched_name': getattr(obj, 'customer_name', ''),
            'changes': changes,
        })

    summary = {
        'create': sum(1 for i in items if i['action'] == 'create'),
        'update': sum(1 for i in items if i['action'] == 'update'),
        'skip': sum(1 for i in items if i['action'] == 'skip'),
        'error': len(errors),
    }
    return {
        'items': items,
        'errors': errors,
        'warnings': warnings,
        'summary': summary,
    }


# ---------------------------------------------------------------- 写入

def apply_rows(rows, overwrite=False):
    """按预览结果写库：先把阻断错误全部拦下，再单事务全成全败。"""
    preview = build_preview(rows, overwrite=overwrite)
    if preview['errors']:
        raise ValueError('导入数据校验未通过：' + '；'.join(preview['errors'][:10]))

    created = updated = skipped = 0
    with transaction.atomic():
        for item in preview['items']:
            values = item['values']
            if item['action'] == 'create':
                obj = CustomerLedger(**values)
                obj._credit_change_source = 'Excel导入'
                obj.save()
                created += 1
                continue
            if item['action'] == 'skip':
                skipped += 1
                continue
            obj = CustomerLedger.objects.select_for_update().get(pk=item['matched_id'])
            touched = False
            for field in IMPORT_FIELDS:
                new = values.get(field)
                if _empty(new):
                    continue
                old = getattr(obj, field)
                if not overwrite and not _empty(old):
                    continue
                if old == new:
                    continue
                setattr(obj, field, new)
                touched = True
            if touched:
                obj._credit_change_source = 'Excel导入'
                obj.save()
                updated += 1
            else:
                skipped += 1
    return {
        'created': created,
        'updated': updated,
        'skipped': skipped,
        'errors': preview['errors'],
        'warnings': preview['warnings'],
    }


# ---------------------------------------------------------------- 导出 / 模板

def _amount_out(text, number):
    if text:
        return text
    if number is None:
        return ''
    return _display(number)


def ledger_export_row(obj, serial):
    """单条记录 -> 31 列导出值（与导入表头严格对齐）。"""
    return [
        serial,
        obj.responsible_salesperson, obj.customer_name, obj.customer_category,
        obj.legal_representative, obj.address, obj.credit_code,
        _amount_out(obj.registered_capital_text, obj.registered_capital),
        _amount_out(obj.paid_in_capital_text, obj.paid_in_capital),
        obj.company_status, obj.business_license, obj.payment_terms,
        _amount_out(obj.credit_limit_text, obj.credit_limit),
        _display(obj.change_date),
        obj.basic_contract, _display(obj.contract_sign_date), obj.contract_issues,
        obj.asset_guarantee, obj.legal_id_card, obj.guarantor_id_card,
        obj.asset_guarantor, obj.guaranteed_asset_detail, obj.asset_mortgaged,
        obj.asset_guarantee_issues, obj.contract_count,
        obj.handling_suggestion, obj.system_info, obj.location,
        obj.archive_no or '', _display(obj.archive_date), obj.remarks,
    ]


def template_response():
    """下载导入模板：31 列表头 + 填写说明工作表（不含示例数据，避免误导入）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from work_record.export_utils import content_disposition

    wb = Workbook()
    ws = wb.active
    ws.title = '客户台账'
    ws.append(CUSTOMER_LEDGER_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='E6F7FF')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for col in range(1, len(CUSTOMER_LEDGER_HEADERS) + 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 18

    notes = wb.create_sheet('填写说明')
    notes.append(['填写说明'])
    notes['A1'].font = Font(bold=True, size=14)
    for line in [
        '1. 仅支持 .xlsx / .xls，单次最多 5000 行；「客户名称」为必填列。',
        '2. 去重顺序：统一社会信用代码 → 档案编号 → 客户名称+法人姓名。',
        '3. 默认仅补空，不覆盖系统已有值；导入向导中可切换为覆盖。',
        '4. 注册资金/实缴资金/额度可填「50万」「0」等原文，系统自动换算为万元。',
        '5. 日期支持 2026-01-31、2026/1/31、20260131；无法识别时留空。',
        '6. 同一文件内统一社会信用代码、档案编号不得重复。',
        '7. 附件（营业执照/法人身份证/基本合同等合并 PDF）不在 Excel 中导入，',
        '   请在记录详情或编辑页上传。',
    ]:
        notes.append([line])
    notes.column_dimensions['A'].width = 80

    from work_record.export_utils import sanitize_xlsx_formulas
    sanitize_xlsx_formulas(wb)
    buf = io.BytesIO()
    wb.save(buf)
    filename = f'客户台账导入模板-{timezone.localtime().strftime("%Y%m%d")}.xlsx'
    response = HttpResponse(
        buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = content_disposition(filename)
    return response
