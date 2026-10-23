"""records 业务模块新增/编辑/删除/导出/列表视图工厂。

各业务模块（装修、客户账期、客户客诉、制度及处罚、核价、每日记录、保险、客诉）
的新增/编辑表单、删除、Excel 导出与列表视图完全同构，仅差异项不同。本模块用
工厂函数统一实现权限校验、POST 限制、操作留痕、AJAX 弹窗双模式、搜索/筛选/
分页、Excel 生成与公式注入防护，业务侧只声明差异项。

注意：URL 配置（records/urls.py）按视图名引用，替换为工厂生成后视图名保持不变。
"""
from datetime import datetime

from django.core.paginator import Paginator
from django.db.models import Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.template.loader import render_to_string
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from accounts.decorators import perm_required

from .audit import log_operation, snapshot


def _sanitize_xlsx_formulas(wb):
    """导出前防护公式注入（实现已收敛到 work_record.export_utils）。"""
    from work_record.export_utils import sanitize_xlsx_formulas

    return sanitize_xlsx_formulas(wb)


def make_form_view(*, model, form_class, perm_module, log_label, display_field,
                   redirect_name, template_name, title_add, title_edit,
                   summary_label=None, modal_template='records/_form_modal.html',
                   after_save=None):
    """生成一对同构新增/编辑视图（整页表单 + AJAX 弹窗双模式，保存成功即留痕）。

    - model/form_class:    目标模型与其 ModelForm；
    - perm_module:         权限模块标识（新增动作为 create，编辑动作为 update）；
    - log_label:           留痕模块名（如「装修」「诉讼」，留痕对象标识前缀同源）；
    - summary_label:       留痕摘要中的业务名称，默认同 log_label
                           （装修留痕摘要为「装修记录」，与模块名不同）；
    - display_field:       留痕摘要中展示的对象字段名；
    - redirect_name:       保存成功后跳转的 URL 名称；
    - template_name:       非弹窗模式的整页表单模板；
    - title_add/title_edit: 整页表单标题（如「新增核价」「编辑核价」）；
    - modal_template:      AJAX 弹窗局部模板（records 各模块统一为 _form_modal.html）；
    - after_save:          可选保存后钩子 after_save(obj, created)，在留痕前调用
                           （如保险保存后自动挂接标的/保单年度），默认 None 零影响。

    AJAX 语义：GET 返回 {'ok': True, 'html': 弹窗}；校验失败返回
    {'ok': False, 'html': 弹窗} 且状态码 422；成功返回 {'ok': True, 'message': ...}。
    非弹窗模式：校验失败时整页重渲染表单并展示错误。
    """
    label = summary_label or log_label

    def _modal(request, form, save_label, ok):
        html = render_to_string(modal_template, {'form': form, 'save_label': save_label},
                                request=request)
        return JsonResponse({'ok': ok, 'html': html}, status=200 if ok else 422)

    @perm_required(perm_module, "create")
    def _add_view(request):
        is_ajax = request.headers.get('x-requested-with') == 'XMLHttpRequest'
        if request.method == 'POST':
            form = form_class(request.POST, request.FILES)
            if form.is_valid():
                obj = form.save()
                if after_save is not None:
                    after_save(obj, True)
                log_operation(request, log_label, '新增', f'[{log_label}]下的记录#{obj.id}',
                              f'新增{label}: {getattr(obj, display_field)}',
                              after=snapshot(obj))
                if is_ajax:
                    return JsonResponse({'ok': True, 'message': '新增成功'})
                return redirect(redirect_name)
            if is_ajax:
                return _modal(request, form, '保存记录', False)
        else:
            form = form_class()
        if is_ajax:
            return _modal(request, form, '保存记录', True)
        return render(request, template_name, {'form': form, 'title': title_add})

    @perm_required(perm_module, "update")
    def _edit_view(request, pk):
        obj = get_object_or_404(model, pk=pk)
        # 快照须在绑定表单前拍：ModelForm.is_valid() 的 _post_clean 会就地改写实例
        before_snap = snapshot(obj)
        is_ajax = request.headers.get('x-requested-with') == 'XMLHttpRequest'
        if request.method == 'POST':
            form = form_class(request.POST, request.FILES, instance=obj)
            if form.is_valid():
                form.save()
                if after_save is not None:
                    after_save(obj, False)
                log_operation(request, log_label, '更新', f'[{log_label}]下的记录#{obj.id}',
                              f'编辑{label}: {getattr(obj, display_field)}',
                              before=before_snap, after=snapshot(obj))
                if is_ajax:
                    return JsonResponse({'ok': True, 'message': '保存成功'})
                return redirect(redirect_name)
            if is_ajax:
                return _modal(request, form, '保存修改', False)
        else:
            form = form_class(instance=obj)
        if is_ajax:
            return _modal(request, form, '保存修改', True)
        return render(request, template_name, {'form': form, 'title': title_edit})

    return _add_view, _edit_view


def make_delete_view(*, model=None, perm_module, label, display_field,
                     redirect_name, getter=None, after_delete=None):
    """生成标准删除视图：仅 POST（防跨站触发删除）+ 权限校验 + 操作留痕。

    - model:        被删模型；getter 未提供时按 get_object_or_404(model, pk=pk) 取数；
    - perm_module:  权限模块标识（同 perm_required 第一参数）；
    - label:        业务名称，用于留痕文案（如「装修」「每日记录」）；
    - display_field:留痕摘要中展示的对象字段名；
    - redirect_name:删除成功后跳转的 URL 名称；
    - getter:       自定义取数 (request, pk) -> obj，用于行级数据隔离
                    （如每日记录必须经 _get_my_record 校验归属，禁止跨用户删除）；
    - after_delete: 可选删除后钩子 after_delete(obj)，在 obj.delete() 之后、
                    留痕之前调用；obj 的 pk 已置 None，其余字段值仍在
                    （如保险删除后按原 policy_year_id 重算/清理保单年度）。
    """
    @perm_required(perm_module, "delete")
    @require_POST
    def _delete_view(request, pk):
        obj = getter(request, pk) if getter else get_object_or_404(model, pk=pk)
        info = f'删除{label}: {getattr(obj, display_field)}'
        obj_id = obj.pk  # delete() 会将实例 pk 置 None，须先取出再删除
        before_snap = snapshot(obj)
        obj.delete()
        if after_delete is not None:
            after_delete(obj)
        log_operation(request, label, '删除', f'[{label}]下的记录#{obj_id}', info,
                      before=before_snap)
        return redirect(redirect_name)

    return _delete_view


def make_export_view(*, model=None, perm_module, sheet_title, filename_base,
                     headers, row_values, queryset=None, column_width=20,
                     empty_text, log_label=None):
    """生成标准 Excel 导出视图：空数据提示 + 列宽 + 公式注入防护。

    - model:        数据模型；queryset 未提供时按 model.objects.all() 导出；
    - sheet_title:  工作表名称；
    - filename_base:下载文件名前缀（自动追加 -YYYYMMDD.xlsx）；
    - headers:      表头列表（含「序号」列时由 row_values 自行返回序号值）；
    - row_values:   (obj, serial) -> list[单元格值]，serial 为从 1 起的记录序号；
    - queryset:     自定义取数 (request) -> QuerySet，用于数据隔离/排序
                    （如每日记录仅导出当前用户的记录）；
    - column_width: 数据列宽（默认 20，个别模块为 30）；
    - empty_text:   空数据提示完整文案（逐字传入，避免拼接歧义）；
    - log_label:    导出留痕的模块名（默认取 sheet_title），写入 operation_type='导出'。
    """
    audit_label = log_label or sheet_title

    @perm_required(perm_module, "read")
    def _export_view(request):
        objs = queryset(request) if queryset else model.objects.all()
        # 归属隔离兜底：含 user 字段的模型（如每日记录）即便未来加入导出清单，
        # 也不会把其他用户的记录导出给普通用户。
        if model is not None:
            has_owner = any(
                getattr(f, 'name', '') == 'user' and getattr(f, 'is_relation', False)
                for f in model._meta.get_fields()
            )
            user = getattr(request, 'user', None)
            if has_owner and user is not None and getattr(user, 'is_authenticated', False) \
                    and not getattr(user, 'is_superuser', False):
                objs = objs.filter(user=user)

        # 提前取总数：迭代后再 count() 会触发额外全表扫描（审计 BE-14）
        total = objs.count()
        if not total:
            return HttpResponse(empty_text)

        wb = Workbook()
        ws = wb.active
        ws.title = sheet_title

        # 表头
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)

        # 数据行：分块迭代，避免 QuerySet 结果集全量驻留内存
        for serial, obj in enumerate(objs.iterator(chunk_size=2000), 1):
            for col, value in enumerate(row_values(obj, serial), 1):
                ws.cell(row=serial + 1, column=col, value=value)

        # 设置列宽
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = column_width

        # 生成文件名（数据库名-YYYYMMDD.xlsx）
        filename = f'{filename_base}-{datetime.now().strftime("%Y%m%d")}.xlsx'
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        from work_record.export_utils import content_disposition

        response['Content-Disposition'] = content_disposition(filename)
        _sanitize_xlsx_formulas(wb)
        wb.save(response)
        # 导出行为留痕（operation_type='导出'）：供审计异常检测识别批量导出
        log_operation(request, audit_label, '导出', f'[{audit_label}]下的导出',
                      f'导出{audit_label}: {total} 条')
        return response

    return _export_view


def _per_page(request):
    """列表页每页条数：仅允许 10/30/50，默认 10"""
    try:
        v = int(request.GET.get('per_page', 10))
    except (TypeError, ValueError):
        v = 10
    return v if v in (10, 30, 50) else 10


def _dedupe_values(qs, field):
    """下拉选项来源：字段值去重（strip）、剔除空白、排序（与原各列表视图一致）。"""
    return sorted({v.strip() for v in qs.values_list(field, flat=True) if v and v.strip()})


def _apply_filters(qs, request, filters):
    """按筛选声明处理 queryset，返回 (qs, context片段)。

    筛选声明（tuple）：
    - (param, 'lookup')           GET[param] 非空时精确/模糊过滤 filter(**{lookup: value})；
    - (param, 'lookup', fn)       值先经 fn(value) 转换再过滤（如 'true' -> True）；
    - (param, callable)           完全自定义过滤 callable(qs, value) -> qs（如到期状态分支）；
    - (param, 'lookup', fn, key)  第 4 元素自定义 context 键名（默认 f'{param}_filter'）。

    transform 传字面量 'strip' 表示 GET 值先去首尾空白，再进入 context 与过滤
    （samples 模块语义：筛选值与回填值均 strip）。

    GET 原始值无论是否为空都会写入 context（模板回填筛选框用）。
    """
    ctx = {}
    for spec in (filters or []):
        if not 2 <= len(spec) <= 4:
            raise ValueError(f'非法筛选声明: {spec!r}')
        param, second = spec[0], spec[1]
        transform = spec[2] if len(spec) > 2 else None
        key = spec[3] if len(spec) > 3 else f'{param}_filter'
        value = request.GET.get(param, '')
        if transform == 'strip':
            value = value.strip()
            transform = None
        ctx[key] = value
        if not value:
            continue
        if callable(second):
            qs = second(qs, value)
        elif transform:
            qs = qs.filter(**{second: transform(value)})
        else:
            qs = qs.filter(**{second: value})
    return qs, ctx


def make_list_view(*, perm_module, template_name, page_var, model=None,
                   queryset=None, search_fields=None, filters=None,
                   filter_options=None, option_queries=None, extra_context=None,
                   search_param='search', per_page=None):
    """生成标准列表视图：全文搜索 + 多条件筛选 + 下拉选项 + 分页。

    - perm_module:    权限模块标识（perm_required 第一参数，动作为 read）；
    - template_name:  列表模板路径；
    - page_var:       分页对象写入 context 的键名（如 'decorations'）；
    - model:          数据模型；queryset 未提供时按 model.objects.all() 取数；
    - queryset:       自定义取数 (request) -> QuerySet，用于数据隔离/排序
                      （如每日记录仅展示当前用户记录并按日期倒序）；
    - search_fields:  全文搜索字段（icontains，Q| 连接，对应搜索 GET 参数）；
    - filters:        筛选声明列表，格式见 _apply_filters；
    - filter_options: {context键: 字段名}，基于 model.objects 全表生成去重选项
                      （不受本次筛选影响，与原实现一致）；
    - option_queries: {context键: (字段名, getter)}，getter(request) -> QuerySet，
                      用于选项来源受数据隔离约束的场景（每日记录部门列表）；
    - extra_context:  额外固定注入 context（如硬编码的 choices）；
    - search_param:   搜索 GET 参数名。'search'（默认）时 context 键为 search_term、
                      值不剥离（records 模块语义）；其他值时 context 键与参数同名、
                      值 strip（samples 模块语义，如 keyword）；
    - per_page:       固定每页条数（如操作留痕固定 20）；None 时按 GET['per_page']
                      在 10/30/50 中取值（默认 10）。
    """

    @perm_required(perm_module, "read")
    def _list_view(request):
        base = queryset(request) if queryset else model.objects.all()

        # 全文搜索（参数名与剥离语义因模块而异，见 search_param）
        search_term = request.GET.get(search_param, '')
        if search_param != 'search':
            search_term = search_term.strip()
            search_key = search_param
        else:
            search_key = 'search_term'
        if search_term and search_fields:
            cond = Q()
            for field in search_fields:
                cond |= Q(**{f'{field}__icontains': search_term})
            base = base.filter(cond)

        # 多条件筛选（GET 原值一并回传模板）
        base, filter_ctx = _apply_filters(base, request, filters)

        # 下拉选项
        options = {}
        for key, field in (filter_options or {}).items():
            options[key] = _dedupe_values(model.objects, field)
        for key, (field, getter) in (option_queries or {}).items():
            options[key] = _dedupe_values(getter(request), field)

        # 分页（固定条数优先，否则按 GET 10/30/50）
        paginator = Paginator(base, per_page if per_page else _per_page(request))
        page = paginator.get_page(request.GET.get('page', 1))

        context = {page_var: page, search_key: search_term, **filter_ctx, **options}
        context.update(extra_context or {})
        return render(request, template_name, context)

    return _list_view
