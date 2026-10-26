"""客户台账模块测试：CRUD、附件、Excel 导入导出、权限与受保护媒体。"""

import io
import json
import shutil
import tempfile
import uuid
from datetime import date
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse

from accounts.models import Permission, Role

from . import customer_ledger_import as ledger_io
from .models import CustomerLedger

User = get_user_model()

PDF_BYTES = b'%PDF-1.4\n1 0 obj\n<<>>\nendobj\ntrailer\n<<>>\n%%EOF'


def grant_ledger(user, *actions):
    role = Role.objects.create(code=f'ledger-{user.username}', name='客户台账测试角色')
    role.permissions.set(
        Permission.objects.filter(code__in=[f'customer_ledger.{a}' for a in actions])
    )
    role.users.add(user)
    cache.clear()
    return role


def build_xlsx(rows):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(ledger_io.CUSTOMER_LEDGER_HEADERS)
    for row in rows:
        ws.append([row.get(header, '') for header in ledger_io.CUSTOMER_LEDGER_HEADERS])
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(
        'ledger.xlsx', buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


class CustomerLedgerTestBase(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        call_command('sync_rbac', verbosity=0)

    def setUp(self):
        self.media_root = tempfile.mkdtemp(prefix='ledger-media-')
        self._override = override_settings(MEDIA_ROOT=self.media_root)
        self._override.enable()
        self.addCleanup(self._override.disable)
        self.addCleanup(shutil.rmtree, self.media_root, ignore_errors=True)

        self.user = User.objects.create_user(username='ledger_user', password='pass12345')
        grant_ledger(self.user, 'read', 'create', 'update', 'delete')
        self.client.force_login(self.user)

    def make_ledger(self, **overrides):
        data = {
            'customer_name': '测试客户',
            'responsible_salesperson': '张三',
            'customer_category': '经销商',
            'credit_code': '91330000TEST0001X',
            'payment_terms': '月结30天',
            'credit_limit_text': '80万',
            'credit_limit': Decimal('80.00'),
            'archive_no': f'A-{uuid.uuid4().hex[:8]}',
        }
        data.update(overrides)
        return CustomerLedger.objects.create(**data)


class CustomerLedgerPermissionTests(CustomerLedgerTestBase):
    def test_list_requires_permission(self):
        other = User.objects.create_user(username='no_ledger', password='pass12345')
        self.client.force_login(other)
        resp = self.client.get(reverse('customer_ledger_list'))
        self.assertNotEqual(resp.status_code, 200)

    def test_sidebar_entry_follows_customer_complaint(self):
        resp = self.client.get(reverse('home'))
        html = resp.content.decode()
        self.assertIn('客户台账', html)
        self.assertLess(html.find('客户客诉'), html.find('客户台账'))


class CustomerLedgerCrudTests(CustomerLedgerTestBase):
    def test_create_with_pdf_attachment(self):
        pdf = SimpleUploadedFile('scan.pdf', PDF_BYTES, content_type='application/pdf')
        resp = self.client.post(reverse('customer_ledger_add'), {
            'customer_name': '新建客户',
            'responsible_salesperson': '李四',
            'credit_code': '91330000TEST0002Y',
            'archive_no': 'A-002',
            'attachment': pdf,
        })
        self.assertEqual(resp.status_code, 302)
        ledger = CustomerLedger.objects.get(customer_name='新建客户')
        self.assertTrue(ledger.has_attachment)
        self.assertTrue(ledger.attachment.name.startswith('customer_ledgers/'))
        self.assertTrue(ledger.attachment.storage.exists(ledger.attachment.name))

    def test_reject_non_pdf_attachment(self):
        fake = SimpleUploadedFile('scan.pdf', b'hello world', content_type='application/pdf')
        resp = self.client.post(reverse('customer_ledger_add'), {
            'customer_name': '伪造附件客户',
            'attachment': fake,
        })
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'PDF')
        self.assertFalse(CustomerLedger.objects.filter(customer_name='伪造附件客户').exists())

    def test_archive_no_must_be_unique(self):
        self.make_ledger(customer_name='客户一', archive_no='DUP-1', credit_code='C1')
        resp = self.client.post(reverse('customer_ledger_add'), {
            'customer_name': '客户二',
            'archive_no': 'DUP-1',
            'credit_code': 'C2',
        })
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, '档案编号')
        self.assertFalse(CustomerLedger.objects.filter(customer_name='客户二').exists())

    def test_blank_archive_no_allows_multiple(self):
        self.make_ledger(customer_name='空编号一', archive_no=None, credit_code='C10')
        self.make_ledger(customer_name='空编号二', archive_no=None, credit_code='C11')
        self.assertEqual(CustomerLedger.objects.filter(archive_no__isnull=True).count(), 2)

    def test_edit_replaces_and_deletes_old_attachment(self):
        ledger = self.make_ledger()
        first = SimpleUploadedFile('first.pdf', PDF_BYTES, content_type='application/pdf')
        self.client.post(reverse('customer_ledger_edit', args=[ledger.pk]), {
            'customer_name': ledger.customer_name,
            'credit_code': ledger.credit_code,
            'archive_no': ledger.archive_no,
            'attachment': first,
        })
        ledger.refresh_from_db()
        old_name = ledger.attachment.name
        self.assertTrue(ledger.attachment.storage.exists(old_name))

        second = SimpleUploadedFile('second.pdf', PDF_BYTES, content_type='application/pdf')
        self.client.post(reverse('customer_ledger_edit', args=[ledger.pk]), {
            'customer_name': ledger.customer_name,
            'credit_code': ledger.credit_code,
            'archive_no': ledger.archive_no,
            'attachment': second,
        })
        ledger.refresh_from_db()
        self.assertNotEqual(ledger.attachment.name, old_name)
        self.assertFalse(ledger.attachment.storage.exists(old_name))

    def test_delete_removes_attachment_file(self):
        ledger = self.make_ledger()
        pdf = SimpleUploadedFile('del.pdf', PDF_BYTES, content_type='application/pdf')
        self.client.post(reverse('customer_ledger_edit', args=[ledger.pk]), {
            'customer_name': ledger.customer_name,
            'credit_code': ledger.credit_code,
            'archive_no': ledger.archive_no,
            'attachment': pdf,
        })
        ledger.refresh_from_db()
        name = ledger.attachment.name
        storage = ledger.attachment.storage
        self.assertTrue(storage.exists(name))

        resp = self.client.post(reverse('customer_ledger_delete', args=[ledger.pk]))
        self.assertEqual(resp.status_code, 302)
        self.assertFalse(CustomerLedger.objects.filter(pk=ledger.pk).exists())
        self.assertFalse(storage.exists(name))

    def test_detail_page_and_modal(self):
        ledger = self.make_ledger()
        resp = self.client.get(reverse('customer_ledger_detail', args=[ledger.pk]))
        self.assertContains(resp, '客户台账详细信息')
        resp = self.client.get(
            reverse('customer_ledger_detail', args=[ledger.pk]),
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertEqual(resp.status_code, 200)
        self.assertIn('客户台账详情', json.loads(resp.content)['html'])

    def test_attachment_endpoint_requires_existing_file(self):
        ledger = self.make_ledger()
        resp = self.client.get(reverse('customer_ledger_attachment', args=[ledger.pk]))
        self.assertEqual(resp.status_code, 404)

    def test_protected_media_requires_module_permission(self):
        ledger = self.make_ledger()
        ledger.attachment = SimpleUploadedFile('p.pdf', PDF_BYTES, content_type='application/pdf')
        ledger.save()
        url = ledger.attachment.url

        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)

        other = User.objects.create_user(username='ledger_noaccess', password='pass12345')
        self.client.force_login(other)
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 403)

    def test_list_search_and_filters(self):
        self.make_ledger(customer_name='杭州甲客户', customer_category='经销商', credit_code='C-A')
        self.make_ledger(customer_name='宁波乙客户', customer_category='终端商', credit_code='C-B')
        resp = self.client.get(reverse('customer_ledger_list'), {'search': '杭州'})
        self.assertContains(resp, '杭州甲客户')
        self.assertNotContains(resp, '宁波乙客户')
        resp = self.client.get(reverse('customer_ledger_list'), {'customer_category': '终端商'})
        self.assertContains(resp, '宁波乙客户')
        self.assertNotContains(resp, '杭州甲客户')

    def test_attachment_filter(self):
        self.make_ledger(customer_name='有附件客户', credit_code='C-ATT')
        target = CustomerLedger.objects.get(customer_name='有附件客户')
        target.attachment = SimpleUploadedFile('a.pdf', PDF_BYTES, content_type='application/pdf')
        target.save()
        self.make_ledger(customer_name='无附件客户', credit_code='C-NOATT')

        resp = self.client.get(reverse('customer_ledger_list'), {'attachment': 'yes'})
        self.assertContains(resp, '有附件客户')
        self.assertNotContains(resp, '无附件客户')
        resp = self.client.get(reverse('customer_ledger_list'), {'attachment': 'no'})
        self.assertContains(resp, '无附件客户')
        self.assertNotContains(resp, '有附件客户')


class CustomerLedgerExportTests(CustomerLedgerTestBase):
    def test_export_headers_match_import_template(self):
        from openpyxl import load_workbook

        self.make_ledger()
        resp = self.client.get(reverse('customer_ledger_export'))
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        self.assertEqual(headers, ledger_io.CUSTOMER_LEDGER_HEADERS)
        self.assertEqual(ws.max_row, 2)

    def test_template_download(self):
        from openpyxl import load_workbook

        resp = self.client.get(reverse('customer_ledger_template'))
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(io.BytesIO(resp.content))
        headers = [cell.value for cell in wb['客户台账'][1]]
        self.assertEqual(headers, ledger_io.CUSTOMER_LEDGER_HEADERS)
        self.assertIn('填写说明', wb.sheetnames)


class CustomerLedgerImportTests(CustomerLedgerTestBase):
    def test_import_preview_then_confirm_creates_and_updates(self):
        existing = self.make_ledger(
            customer_name='已有客户', credit_code='91330000EXIST0001',
            responsible_salesperson='',
        )
        upload = build_xlsx([
            {
                '责任业务': '王五', '客户名称': '新客户', '客户分类': '经销商',
                '统一社会信用代码': '91330000NEW00001', '注册资金': '50万',
                '额度': '30万', '变更时间': '2026-01-15', '档案编号': 'N-001',
            },
            {
                '责任业务': '赵六', '客户名称': '已有客户',
                '统一社会信用代码': '91330000EXIST0001', '实缴资金': '20万',
            },
        ])
        resp = self.client.post(reverse('customer_ledger_import'), {'file': upload})
        self.assertEqual(resp.status_code, 200)
        preview = resp.context['preview']
        self.assertEqual(preview['summary']['create'], 1)
        self.assertEqual(preview['summary']['update'], 1)
        self.assertFalse(preview['errors'])

        payload = resp.context['payload']
        confirm = self.client.post(reverse('customer_ledger_import_confirm'), {
            'payload': payload, 'overwrite': '0',
        })
        self.assertEqual(confirm.status_code, 200)
        self.assertEqual(confirm.context['result']['created'], 1)
        self.assertEqual(confirm.context['result']['updated'], 1)

        created = CustomerLedger.objects.get(customer_name='新客户')
        self.assertEqual(created.registered_capital, Decimal('50.00'))
        self.assertEqual(created.registered_capital_text, '50万')
        self.assertEqual(created.credit_limit, Decimal('30.00'))
        self.assertEqual(created.change_date, date(2026, 1, 15))
        existing.refresh_from_db()
        self.assertEqual(existing.responsible_salesperson, '赵六')
        self.assertEqual(existing.paid_in_capital, Decimal('20.00'))

    def test_only_fill_mode_does_not_overwrite_existing(self):
        self.make_ledger(
            customer_name='保留客户', credit_code='91330000KEEP00001',
            responsible_salesperson='原业务',
        )
        upload = build_xlsx([{
            '客户名称': '保留客户', '统一社会信用代码': '91330000KEEP00001',
            '责任业务': '新业务',
        }])
        resp = self.client.post(reverse('customer_ledger_import'), {'file': upload})
        preview = resp.context['preview']
        self.assertEqual(preview['summary']['skip'], 1)
        self.client.post(reverse('customer_ledger_import_confirm'), {
            'payload': resp.context['payload'], 'overwrite': '0',
        })
        obj = CustomerLedger.objects.get(customer_name='保留客户')
        self.assertEqual(obj.responsible_salesperson, '原业务')

    def test_overwrite_mode_updates_existing(self):
        self.make_ledger(
            customer_name='覆盖客户', credit_code='91330000OVER00001',
            responsible_salesperson='原业务',
        )
        upload = build_xlsx([{
            '客户名称': '覆盖客户', '统一社会信用代码': '91330000OVER00001',
            '责任业务': '新业务',
        }])
        resp = self.client.post(reverse('customer_ledger_import'), {
            'file': upload, 'overwrite': '1',
        })
        self.assertEqual(resp.context['preview']['summary']['update'], 1)
        self.client.post(reverse('customer_ledger_import_confirm'), {
            'payload': resp.context['payload'], 'overwrite': '1',
        })
        obj = CustomerLedger.objects.get(customer_name='覆盖客户')
        self.assertEqual(obj.responsible_salesperson, '新业务')

    def test_duplicate_credit_code_blocks_import(self):
        upload = build_xlsx([
            {'客户名称': '重复一', '统一社会信用代码': '91330000DUP000001'},
            {'客户名称': '重复二', '统一社会信用代码': '91330000DUP000001'},
        ])
        resp = self.client.post(reverse('customer_ledger_import'), {'file': upload})
        self.assertTrue(resp.context['preview']['errors'])
        confirm = self.client.post(reverse('customer_ledger_import_confirm'), {
            'payload': resp.context['payload'], 'overwrite': '0',
        })
        self.assertEqual(confirm.status_code, 200)
        self.assertContains(confirm, '统一社会信用代码')
        self.assertFalse(CustomerLedger.objects.filter(customer_name__startswith='重复').exists())

    def test_reject_non_excel_upload(self):
        bad = SimpleUploadedFile('data.csv', b'customer_name\nfoo\n', content_type='text/csv')
        resp = self.client.post(reverse('customer_ledger_import'), {'file': bad})
        self.assertContains(resp, '仅支持')
        self.assertFalse(CustomerLedger.objects.exists())


class CustomerLedgerAmountParsingTests(TestCase):
    def test_amount_normalization(self):
        cases = {
            '50万': Decimal('50.00'),
            '0': Decimal('0.00'),
            '500000元': Decimal('50.00'),
            '/': None,
            '实缴未披露/查不到': None,
            '': None,
        }
        for raw, expected in cases.items():
            self.assertEqual(ledger_io._to_decimal(raw), expected, raw)

    def test_date_parsing_variants(self):
        self.assertEqual(ledger_io._to_date('2026-01-31'), date(2026, 1, 31))
        self.assertEqual(ledger_io._to_date('2026/1/31'), date(2026, 1, 31))
        self.assertEqual(ledger_io._to_date('20260131'), date(2026, 1, 31))
        self.assertIsNone(ledger_io._to_date(''))
        self.assertIsNone(ledger_io._to_date('12201'))
