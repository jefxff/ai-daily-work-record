"""样品台账 Excel 导出（openpyxl 多 sheet）。

导出范围与列表筛选条件完全一致；中文文件名按 RFC 5987 做 URL 编码，
避免 Content-Disposition 直接携带中文导致的下载中断问题。
"""

from datetime import date

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from urllib.parse import quote

HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
HEADER_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="center")


def _sanitize_xlsx_formulas(wb):
    """导出前防护公式注入（实现已收敛到 work_record.export_utils）。"""
    from work_record.export_utils import sanitize_xlsx_formulas

    return sanitize_xlsx_formulas(wb)


def _write_sheet(ws, headers, rows):
    ws.append(headers)
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="center")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    # 列宽：按内容粗略自适应（中文按 2 字宽）
    for col_idx, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx, values_only=True):
            value = row[0]
            if value is not None:
                text = str(value)
                width = sum(2 if ord(ch) > 127 else 1 for ch in text)
                max_len = max(max_len, min(width, 60))
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 4
    # 自动换行的数据区
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP


def _yn(value):
    return "是" if value else "否"


def export_samples_xlsx(filter_qs):
    """按筛选后的 Sample queryset 导出三个 sheet。"""
    wb = Workbook()

    # Sheet1 样品主档
    ws1 = wb.active
    ws1.title = "样品主档"
    headers1 = [
        "样品编号", "物料名称", "型号", "物料规格", "来料数量",
        "供应商", "送样时间", "当前状态", "登记人", "登记时间",
        "试产轮次", "最新试样结论", "小批量批次", "最新小批量结果",
        "是否已录入金蝶", "终止原因",
    ]
    rows1 = []
    for s in filter_qs.select_related("supplier"):
        latest_test = s.test_records.first()
        latest_batch = s.batch_records.first()
        kingdee = any(r.kingdee_entered for r in s.batch_records.all())
        rows1.append([
            s.sample_no,
            s.material_name,
            s.model,
            s.spec,
            s.incoming_qty,
            s.supplier.name if s.supplier else "",
            s.sample_date.strftime("%Y-%m-%d"),
            s.get_status_display(),
            s.registrar or "",
            s.registered_at.strftime("%Y-%m-%d %H:%M"),
            s.latest_test_round,
            latest_test.get_review_conclusion_display() if latest_test else "",
            s.latest_batch_round,
            latest_batch.get_batch_result_display() if latest_batch else "",
            _yn(kingdee),
            s.terminate_reason,
        ])
    _write_sheet(ws1, headers1, rows1)

    # Sheet2 试样记录（按筛选样品范围）
    ws2 = wb.create_sheet("试样记录")
    headers2 = [
        "试样编号", "样品编号", "物料名称", "型号", "试产轮次", "试产日期",
        "测试结果", "评审结论", "验证报告编码", "是否通知责任人",
        "封装形式", "来料日期", "操作人", "操作时间", "备注",
    ]
    rows2 = []
    tests = (
        filter_qs.prefetch_related("test_records")
        .values_list(
            "sample_no", "material_name", "model",
            "test_records__test_no", "test_records__round_no",
            "test_records__test_date", "test_records__test_result",
            "test_records__review_conclusion", "test_records__report_code",
            "test_records__notified", "test_records__package_type",
            "test_records__incoming_date", "test_records__operator",
            "test_records__operated_at", "test_records__remark",
        )
        .order_by("sample_no", "-test_records__round_no")
    )
    from .models import TestRecord
    conclusion_display = dict(TestRecord.Conclusion.choices)
    for (
        sample_no, material, model, test_no, round_no, test_date, test_result,
        conclusion, report_code, notified, package_type, incoming_date,
        operator, operated_at, remark,
    ) in tests:
        if not test_no:
            continue
        rows2.append([
            test_no, sample_no, material, model, f"第{round_no}轮",
            test_date.strftime("%Y-%m-%d") if test_date else "",
            test_result, conclusion_display.get(conclusion, conclusion or ""),
            report_code, _yn(notified), package_type,
            incoming_date.strftime("%Y-%m-%d") if incoming_date else "",
            operator or "",
            operated_at.strftime("%Y-%m-%d %H:%M") if operated_at else "",
            remark,
        ])
    _write_sheet(ws2, headers2, rows2)

    # Sheet3 小批量记录
    ws3 = wb.create_sheet("小批量记录")
    headers3 = [
        "小批量编号", "样品编号", "物料名称", "型号", "批次轮次",
        "是否下单", "小批量结果", "试样单价", "可替代型号",
        "是否录入金蝶", "操作人", "操作时间", "批次备注",
    ]
    rows3 = []
    from .models import BatchRecord
    result_display = dict(BatchRecord.Result.choices)
    batches = (
        filter_qs.prefetch_related("batch_records")
        .values_list(
            "sample_no", "material_name", "model",
            "batch_records__batch_no", "batch_records__round_no",
            "batch_records__ordered", "batch_records__batch_result",
            "batch_records__unit_price", "batch_records__alternative_model",
            "batch_records__kingdee_entered", "batch_records__operator",
            "batch_records__operated_at", "batch_records__remark",
        )
        .order_by("sample_no", "-batch_records__round_no")
    )
    for (
        sample_no, material, model, batch_no, round_no, ordered, result,
        unit_price, alt_model, kingdee, operator, operated_at, remark,
    ) in batches:
        if not batch_no:
            continue
        rows3.append([
            batch_no, sample_no, material, model, f"第{round_no}批",
            _yn(ordered), result_display.get(result, result or ""),
            float(unit_price) if unit_price is not None else "",
            alt_model, _yn(kingdee), operator or "",
            operated_at.strftime("%Y-%m-%d %H:%M") if operated_at else "",
            remark,
        ])
    _write_sheet(ws3, headers3, rows3)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"试样跟进台账_{date.today():%Y%m%d}.xlsx"
    # RFC 5987：filename* 携带 URL 编码中文，兼容主流浏览器
    from work_record.export_utils import content_disposition

    response["Content-Disposition"] = content_disposition(filename)
    _sanitize_xlsx_formulas(wb)
    wb.save(response)
    return response
